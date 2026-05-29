from io import BytesIO

from django.contrib.auth import get_user_model
from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook

from .forms import OrderForm
from .models import Customer, Order


class OrderFormTests(TestCase):
    def test_quantity_must_be_at_least_one(self):
        form = OrderForm(data={
            "receiver_name": "홍길동",
            "receiver_phone": "010-0000-0000",
            "receiver_address": "서울",
            "product": Order.PRODUCT_5,
            "quantity": 0,
        })

        self.assertFalse(form.is_valid())
        self.assertIn("quantity", form.errors)


class OrderViewTests(TestCase):
    def test_order_new_saves_valid_order(self):
        customer = Customer.objects.create(
            name="홍길동",
            phone="010-0000-0000",
            address="서울",
        )
        session = self.client.session
        session["customer_id"] = customer.id
        session.save()

        response = self.client.post(reverse("order_new"), data={
            "receiver_name": "김철수",
            "receiver_phone": "010-1111-2222",
            "receiver_address": "부산",
            "product": Order.PRODUCT_10,
            "quantity": 2,
        })

        self.assertRedirects(response, reverse("order_done"))
        self.assertEqual(Order.objects.count(), 1)
        self.assertEqual(Order.objects.get().total_price, 120000)


class ExportOrdersTests(TestCase):
    def test_admin_export_escapes_formula_like_values(self):
        user = get_user_model().objects.create_user(
            username="staff",
            password="password",
            is_staff=True,
        )
        customer = Customer.objects.create(
            name="=cmd",
            phone="010-0000-0000",
            address="+서울",
        )
        Order.objects.create(
            customer=customer,
            receiver_name="@수령인",
            receiver_phone="010-1111-2222",
            receiver_address="-주소",
            product=Order.PRODUCT_5,
            quantity=1,
        )

        self.client.force_login(user)
        response = self.client.get(reverse("export_orders_xlsx"))

        self.assertEqual(response.status_code, 200)

        workbook = load_workbook(BytesIO(response.content))
        sheet = workbook.active
        row = next(sheet.iter_rows(min_row=2, max_row=2, values_only=True))

        self.assertEqual(row[1], "'=cmd")
        self.assertEqual(row[5], "'-주소")
